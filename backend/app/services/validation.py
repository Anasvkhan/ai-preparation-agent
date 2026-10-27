"""Upload validation. Every check looks at the actual bytes, not only the file extension."""
import io
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

import pymupdf as fitz

MB = 1024 * 1024
MAX_DOC_BYTES = 25 * MB
MAX_CODE_BYTES = 1 * MB
MAX_CODE_FILES = 5
MIN_PDF_TEXT_CHARS = 200

CODE_EXTENSIONS = {".py", ".js", ".jsx", ".ts", ".tsx", ".html", ".css", ".json", ".sql"}
KINDS = {"report", "presentation", "code"}

PDF_MIME = "application/pdf"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


class ValidationError(Exception):
    pass


@dataclass
class ValidatedFile:
    filename: str
    content_type: str
    file_format: str  # pdf | pptx | text


def _ext(filename: str) -> str:
    return PurePath(filename).suffix.lower()


def _safe_filename(filename: str) -> str:
    name = PurePath(filename.replace("\\", "/")).name.strip()
    if not name or name in {".", ".."}:
        raise ValidationError("Invalid file name")
    return name[:150]


def _check_pdf(data: bytes) -> None:
    if not data.startswith(b"%PDF"):
        raise ValidationError("File is not a valid PDF")
    try:
        with fitz.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass:
                raise ValidationError("Password-protected PDFs are not supported")
            text_chars = sum(len(page.get_text("text").strip()) for page in doc)
    except ValidationError:
        raise
    except Exception:
        raise ValidationError("PDF could not be opened (file may be corrupted)")
    if text_chars < MIN_PDF_TEXT_CHARS:
        raise ValidationError("PDF has no readable text (scanned or image-only PDFs are not supported)")


def _check_pptx(data: bytes) -> None:
    if not data.startswith(b"PK"):
        raise ValidationError("File is not a valid PPTX presentation")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if "ppt/presentation.xml" not in z.namelist():
                raise ValidationError("File is not a valid PPTX presentation")
        from pptx import Presentation

        prs = Presentation(io.BytesIO(data))
        has_text = any(shape.has_text_frame and shape.text_frame.text.strip()
                       for slide in prs.slides for shape in slide.shapes)
    except ValidationError:
        raise
    except Exception:
        raise ValidationError("Presentation could not be opened (file may be corrupted)")
    if not has_text:
        raise ValidationError("Presentation has no readable text")


def _check_code(data: bytes) -> None:
    if b"\x00" in data:
        raise ValidationError("Binary files are not allowed as code evidence")
    try:
        text = data.decode("utf-8", errors="strict")
    except UnicodeDecodeError:
        raise ValidationError("Code file must be valid UTF-8 text")
    if not text.strip():
        raise ValidationError("Code file is empty")


def validate_upload(kind: str, filename: str, data: bytes) -> ValidatedFile:
    if kind not in KINDS:
        raise ValidationError("Unknown file type; choose report, presentation or code")
    name = _safe_filename(filename)
    ext = _ext(name)
    size = len(data)
    if size == 0:
        raise ValidationError("File is empty")

    if kind == "code":
        if ext not in CODE_EXTENSIONS:
            raise ValidationError(f"Code files must be one of: {', '.join(sorted(CODE_EXTENSIONS))}")
        if size > MAX_CODE_BYTES:
            raise ValidationError("Code files must be 1 MB or smaller")
        _check_code(data)
        return ValidatedFile(name, "text/plain; charset=utf-8", "text")

    if size > MAX_DOC_BYTES:
        raise ValidationError("Reports and presentations must be 25 MB or smaller")

    if kind == "report":
        if ext != ".pdf":
            raise ValidationError("Reports must be text-readable PDF files")
        _check_pdf(data)
        return ValidatedFile(name, PDF_MIME, "pdf")

    # presentation
    if ext == ".pdf":
        _check_pdf(data)
        return ValidatedFile(name, PDF_MIME, "pdf")
    if ext == ".pptx":
        _check_pptx(data)
        return ValidatedFile(name, PPTX_MIME, "pptx")
    raise ValidationError("Presentations must be PPTX or text-readable PDF files")

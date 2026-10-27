"""Generate realistic sample FYP evidence files into backend/samples/.

Run: .venv/bin/python -m scripts.make_samples
Used for automated tests, the embedding trial and the demo.
"""
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "samples"

REPORT = [
    ("1 Introduction", [
        "Manual attendance in university classrooms wastes around ten minutes of every lecture and is easy to "
        "manipulate through proxy attendance, where one student marks another student as present.",
        "This project, the Smart Attendance System, automates attendance using face recognition. A classroom "
        "camera captures frames, detects faces, matches them against enrolled students and records attendance "
        "in a central database that teachers can review from a web dashboard.",
        "The main objectives are: to reduce attendance time to under one minute, to eliminate proxy attendance, "
        "to give teachers real-time reports, and to keep student biometric data secure.",
    ]),
    ("2 Literature Review", [
        "Earlier systems used RFID cards, but cards can be shared between students so proxy attendance remains "
        "possible. Fingerprint scanners remove sharing but create queues because only one student can be "
        "scanned at a time.",
        "Face recognition approaches based on Haar cascades were fast but inaccurate under poor lighting. Deep "
        "learning models such as FaceNet map each face to a 128-dimensional embedding where faces of the same "
        "person are close together, which gives far better accuracy in real classrooms.",
        "We selected a FaceNet-style embedding approach with MTCNN for detection because published results show "
        "above 99 percent accuracy on the LFW benchmark while running in real time on a mid-range GPU.",
    ]),
    ("3 System Architecture", [
        "The system follows a three-tier architecture. The capture tier runs on a classroom computer connected "
        "to a 1080p camera. The application tier is a Python Flask REST API that performs recognition. The data "
        "tier is a MySQL database that stores students, courses, sessions and attendance records.",
        "The capture client sends one frame every two seconds to the API. The API detects faces with MTCNN, "
        "computes embeddings, and compares each embedding with enrolled embeddings using cosine similarity. A "
        "match is accepted when similarity is above a threshold of 0.75.",
        "A student is marked present only after being recognised in at least three frames during the first "
        "fifteen minutes of the lecture, which prevents a photo held up briefly from counting as attendance.",
    ]),
    ("4 Database Design", [
        "The database has five main tables: students, courses, enrollments, lecture_sessions and attendance. "
        "The attendance table references both the student and the lecture session and stores the first seen "
        "time and the number of frames in which the student was recognised.",
        "Face embeddings are stored in a separate face_embeddings table encrypted with AES-256, and only the API "
        "service account can read them. Raw face images are deleted after enrollment so that only embeddings "
        "are kept, which reduces privacy risk.",
        "Indexes on student_id and session_id keep attendance report queries under 100 milliseconds for a "
        "semester of data.",
    ]),
    ("5 Face Recognition Module", [
        "Enrollment captures twenty images of each student from different angles. Each image is aligned, "
        "resized to 160 by 160 pixels and passed through the embedding model. The average embedding is stored "
        "as the student's template.",
        "During recognition, each detected face is aligned the same way. Liveness detection checks eye blinking "
        "across consecutive frames to reject printed photos and phone screens.",
        "We tuned the similarity threshold using a validation set of 40 students. A threshold of 0.75 gave a "
        "false acceptance rate of 0.5 percent and a false rejection rate of 3 percent.",
    ]),
    ("6 Testing and Results", [
        "Unit tests cover the API endpoints and the matching logic. Integration tests replay recorded classroom "
        "videos through the full pipeline.",
        "In a pilot with three classes of 45 students over two weeks, the system marked attendance in an average "
        "of 40 seconds per lecture with 96.8 percent accuracy compared with manual roll calls.",
        "Most errors came from students wearing masks or sitting at the back row in poor lighting. Adding a "
        "second camera at the back of the room reduced these errors by half.",
    ]),
    ("7 Conclusion and Future Work", [
        "The Smart Attendance System reduced attendance time from about ten minutes to under one minute and "
        "removed proxy attendance in the pilot classes.",
        "Future work includes masked face recognition, a mobile app for teachers, and integration with the "
        "university learning management system.",
    ]),
]

SLIDES = [
    ("Smart Attendance System", "Final Year Project\nFace recognition based classroom attendance"),
    ("Problem Statement", "Manual roll call wastes ~10 minutes per lecture\nProxy attendance is common\nNo real-time reports"),
    ("Proposed Solution", "Classroom camera + face recognition API\nAutomatic attendance after 3 recognised frames\nTeacher web dashboard"),
    ("Architecture", "Capture client -> Flask REST API -> MySQL\nMTCNN detection, FaceNet embeddings\nCosine similarity threshold 0.75"),
    ("Security and Privacy", "Embeddings encrypted with AES-256\nRaw images deleted after enrollment\nLiveness detection with blink check"),
    ("Results", "Pilot: 3 classes, 45 students, 2 weeks\nAccuracy 96.8%\nAverage 40 seconds per lecture"),
    ("Future Work", "Masked face recognition\nTeacher mobile app\nLMS integration"),
]

CODE = '''"""Face matching service for the Smart Attendance System."""
import numpy as np

SIMILARITY_THRESHOLD = 0.75
MIN_FRAMES_FOR_PRESENT = 3


def cosine_similarity(a: np.ndarray, b: np.ndarray) -> float:
    return float(np.dot(a, b) / (np.linalg.norm(a) * np.linalg.norm(b)))


def match_face(embedding, enrolled):
    """Return the student id with the highest similarity above the threshold, else None."""
    best_id, best_score = None, SIMILARITY_THRESHOLD
    for student_id, template in enrolled.items():
        score = cosine_similarity(embedding, template)
        if score > best_score:
            best_id, best_score = student_id, score
    return best_id


def update_attendance(counts, student_id):
    """Count recognised frames; a student is present after MIN_FRAMES_FOR_PRESENT frames."""
    counts[student_id] = counts.get(student_id, 0) + 1
    return counts[student_id] >= MIN_FRAMES_FOR_PRESENT
'''


def make_report(path: Path) -> None:
    doc = pymupdf.open()
    for heading, paragraphs in REPORT:
        page = doc.new_page()
        y = 72
        page.insert_text((72, y), heading, fontsize=18, fontname="helv")
        y += 36
        for para in paragraphs:
            rect = pymupdf.Rect(72, y, 523, y + 160)
            page.insert_textbox(rect, para, fontsize=11, fontname="helv")
            y += 20 + 14 * (len(para) // 85 + 1)
    doc.save(path)


def make_slides(path: Path) -> None:
    prs = Presentation()
    for i, (title, body) in enumerate(SLIDES):
        slide = prs.slides.add_slide(prs.slide_layouts[0 if i == 0 else 1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        slide.notes_slide.notes_text_frame.text = f"Explain {title.lower()} clearly to the examiner."
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(24)
    prs.save(path)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    make_report(OUT / "smart_attendance_report.pdf")
    make_slides(OUT / "smart_attendance_slides.pptx")
    (OUT / "face_matcher.py").write_text(CODE)
    (OUT / "scanned_image.pdf").write_bytes(_image_only_pdf())
    (OUT / "binary.py").write_bytes(b"\x00\x01\x02 not text")
    print(f"Samples written to {OUT}")


def _image_only_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.draw_rect(pymupdf.Rect(72, 72, 300, 300), color=(0, 0, 0), fill=(0.5, 0.5, 0.5))
    return doc.tobytes()


if __name__ == "__main__":
    main()

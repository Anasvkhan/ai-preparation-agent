"""Turn uploaded files into text segments that carry their exact source location.

A segment never spans two pages / slides, so every chunk made from it can be cited precisely.
"""
import io
import re
from dataclasses import dataclass, field
from statistics import median

import pymupdf as fitz
from pptx import Presentation

CODE_WINDOW_LINES = 60
CODE_OVERLAP_LINES = 10


@dataclass
class Segment:
    text: str
    location: dict
    section: str
    extra: dict = field(default_factory=dict)


def _clean(text: str) -> str:
    text = text.replace("­", "")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# ---------------- PDF ----------------
def _looks_like_heading(text: str, size: float, bold: bool, body_size: float) -> bool:
    words = text.split()
    if not (1 <= len(words) <= 12) or text.endswith((".", ",", ";")):
        return False
    return size >= body_size * 1.15 or (bold and size >= body_size and len(words) <= 8)


def extract_pdf(data: bytes) -> tuple[list[Segment], dict]:
    segments: list[Segment] = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        sizes = [span["size"] for page in doc for block in page.get_text("dict")["blocks"]
                 for line in block.get("lines", []) for span in line["spans"] if span["text"].strip()]
        body_size = median(sizes) if sizes else 11.0
        section = "Introduction"

        for page_no, page in enumerate(doc, start=1):
            buffer: list[str] = []
            buffer_section = section

            def flush():
                text = _clean("\n".join(buffer))
                if text:
                    segments.append(Segment(text, {"page": page_no}, buffer_section))
                buffer.clear()

            for block in page.get_text("dict")["blocks"]:
                if block.get("type") != 0:
                    continue  # image block
                for line in block["lines"]:
                    spans = [s for s in line["spans"] if s["text"].strip()]
                    if not spans:
                        continue
                    text = " ".join(s["text"].strip() for s in spans)
                    size = max(s["size"] for s in spans)
                    bold = all(s["flags"] & 16 for s in spans)
                    if _looks_like_heading(text, size, bold, body_size):
                        flush()
                        section = re.sub(r"^\d+(\.\d+)*\s*", "", text).strip() or text
                        buffer_section = section
                    else:
                        buffer.append(text)
            flush()
        meta = {"page_count": doc.page_count}
    return segments, meta


# ---------------- PPTX ----------------
def extract_pptx(data: bytes) -> tuple[list[Segment], dict]:
    prs = Presentation(io.BytesIO(data))
    segments: list[Segment] = []
    for slide_no, slide in enumerate(prs.slides, start=1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        parts: list[str] = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
        body = _clean("\n".join(parts))
        text = _clean(f"{title}\n{body}") if title else body
        if text:
            segments.append(Segment(text, {"slide": slide_no, "title": title}, title or f"Slide {slide_no}"))
    return segments, {"page_count": len(prs.slides)}


# ---------------- Code ----------------
def extract_code(data: bytes, filename: str) -> tuple[list[Segment], dict]:
    """Code is only decoded as text, never executed."""
    lines = data.decode("utf-8").splitlines()
    segments: list[Segment] = []
    start = 0
    while start < len(lines):
        end = min(start + CODE_WINDOW_LINES, len(lines))
        text = "\n".join(lines[start:end]).strip()
        if text:
            segments.append(Segment(
                f"File {filename} lines {start + 1}-{end}:\n{text}",
                {"file": filename, "line_start": start + 1, "line_end": end},
                filename,
            ))
        if end == len(lines):
            break
        start = end - CODE_OVERLAP_LINES
    return segments, {"line_count": len(lines)}


def extract(file_format: str, data: bytes, filename: str) -> tuple[list[Segment], dict]:
    if file_format == "pdf":
        return extract_pdf(data)
    if file_format == "pptx":
        return extract_pptx(data)
    if file_format == "text":
        return extract_code(data, filename)
    raise ValueError(f"unsupported format {file_format}")
